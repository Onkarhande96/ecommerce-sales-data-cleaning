import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def clean_ecommerce_data():
    raw_path = "raw_sales_data.csv"
    output_path = "sales_report_cleaned.xlsx"
    
    # 1. Load raw data
    raw_df = pd.read_csv(raw_path)
    
    # 2. Data Cleaning pipeline
    clean_df = raw_df.copy()
    
    # Drop duplicates
    clean_df = clean_df.drop_duplicates(subset=["Order_ID"]).reset_index(drop=True)
    
    # Trim whitespace and convert to Proper/Title Case
    clean_df["Customer_Name"] = clean_df["Customer_Name"].astype(str).str.strip().str.title()
    clean_df["City"] = clean_df["City"].astype(str).str.strip().str.title()
    clean_df["Product"] = clean_df["Product"].astype(str).str.strip()
    
    # Clean Unit_Price: remove '$', ',', and convert to numeric
    clean_df["Unit_Price"] = (
        clean_df["Unit_Price"]
        .astype(str)
        .str.replace("$", "", regex=False)
        .str.replace(",", "", regex=False)
        .str.strip()
        .astype(float)
    )
    
    # Standardize Dates
    clean_df["Order_Date"] = pd.to_datetime(clean_df["Order_Date"], format="mixed").dt.strftime("%Y-%m-%d")
    
    # Calculate Total_Amount
    clean_df["Total_Amount"] = clean_df["Quantity"] * clean_df["Unit_Price"]
    
    # 3. Create Summary Tables
    # Product Summary
    product_summary = clean_df.groupby("Product").agg(
        Orders_Count=("Order_ID", "count"),
        Total_Units=("Quantity", "sum"),
        Total_Revenue=("Total_Amount", "sum")
    ).reset_index().sort_values(by="Total_Revenue", ascending=False)
    
    # City Summary
    city_summary = clean_df.groupby("City").agg(
        Orders_Count=("Order_ID", "count"),
        Total_Revenue=("Total_Amount", "sum")
    ).reset_index().sort_values(by="Total_Revenue", ascending=False)
    
    # Overall KPIs
    kpi_data = {
        "Metric": ["Total Cleaned Orders", "Total Units Sold", "Total Revenue (₹ / $)", "Average Order Value"],
        "Value": [
            len(clean_df),
            int(clean_df["Quantity"].sum()),
            float(clean_df["Total_Amount"].sum()),
            round(float(clean_df["Total_Amount"].mean()), 2)
        ]
    }
    kpi_df = pd.DataFrame(kpi_data)
    
    # 4. Save to Excel with openpyxl styling
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        raw_df.to_excel(writer, sheet_name="Raw_Data", index=False)
        clean_df.to_excel(writer, sheet_name="Clean_Data", index=False)
        kpi_df.to_excel(writer, sheet_name="Executive_Summary", startrow=1, startcol=0, index=False)
        product_summary.to_excel(writer, sheet_name="Executive_Summary", startrow=8, startcol=0, index=False)
        city_summary.to_excel(writer, sheet_name="Executive_Summary", startrow=8, startcol=5, index=False)
        
    # Apply professional formatting
    wb = openpyxl.load_workbook(output_path)
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    accent_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    accent_font = Font(name="Calibri", size=11, bold=True, color="000000")
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.views.sheetView[0].showGridLines = True
        
        # Style tables
        for row in ws.iter_rows():
            for cell in row:
                if cell.row == 1 and sheet_name != "Executive_Summary":
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif sheet_name == "Executive_Summary" and cell.row in [2, 9]:
                    cell.fill = accent_fill
                    cell.font = accent_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    
        # Auto-adjust column width
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    wb.save(output_path)
    print(f"Data cleanup complete! Saved to {output_path}")

if __name__ == "__main__":
    clean_ecommerce_data()
